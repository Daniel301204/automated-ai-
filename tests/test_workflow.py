import hashlib
import json
import os
from pathlib import Path
from concurrent.futures import ThreadPoolExecutor
from unittest.mock import Mock

import pytest
from sqlalchemy import select
from workflow import Store, approval_digest, evidence_profile, send_gmail
from live_sources import fetch_public, location_assessment


@pytest.fixture
def store(tmp_path):
    value = Store('sqlite:///' + str(tmp_path / 'test.db'))
    for tenant in ('alice', 'bob'):
        value.ensure_user(tenant)
        value.save_profile(tenant, resume_text='Verified resume evidence', profile=json.dumps({'threshold': 80, 'daily_limit': 2}))
    return value


def prepared(store, tenant='alice', job_id='job1'):
    job = dict(job_id=job_id, title='Verified resume evidence', company='Employer', match_score=90,
               url='https://example.org/job', description='Verified resume evidence')
    store.remember(tenant, [job])
    app = dict(to='jobs@example.org', subject='Application', body='Reviewed message',
               requirements_verified=True, profile_verified=True, contact_verified=True, score=90,
               resume_digest=hashlib.sha256(b'Verified resume evidence').hexdigest())
    pdf = b'%PDF-test-fixture'
    store.save_profile(tenant, resume=pdf)
    store.prepare(tenant, job_id, app, pdf)
    return approval_digest(app['to'], app['subject'], app['body'], pdf)


def test_tenant_isolation(store):
    prepared(store)
    assert store.list_jobs('bob') == []
    assert store.profile('bob')['resume_text'] == 'Verified resume evidence'
    with pytest.raises(Exception):
        store.prepare('bob', 'job1', {}, b'%PDF-fixture')
    assert store.list_jobs('alice')[0]['status'] == 'needs_review'


def test_send_requires_exact_approval_and_is_not_duplicated(store):
    digest = prepared(store)
    sender = Mock(return_value='gmail-id')
    with pytest.raises(ValueError, match='Approval'):
        store.send_approved('alice', 'job1', 'wrong-digest', {}, sender=sender)
    sender.assert_not_called()
    assert store.send_approved('alice', 'job1', digest, {}, sender=sender) == 'gmail-id'
    assert store.list_jobs('alice')[0]['status'] == 'awaiting_response'
    with pytest.raises(ValueError):
        store.send_approved('alice', 'job1', digest, {}, sender=sender)
    sender.assert_called_once()


def test_uncertain_delivery_never_retries(store):
    digest = prepared(store)
    sender = Mock(side_effect=TimeoutError('sensitive provider details'))
    with pytest.raises(RuntimeError, match='Check Gmail Sent'):
        store.send_approved('alice', 'job1', digest, {}, sender=sender)
    assert store.list_jobs('alice')[0]['status'] == 'delivery_unknown'
    with pytest.raises(ValueError):
        store.send_approved('alice', 'job1', digest, {}, sender=sender)
    with pytest.raises(ValueError):
        store.track('alice', 'job1', 'needs_review')
    with pytest.raises(ValueError):
        prepared(store)
    sender.assert_called_once()


def test_daily_limit_is_per_user_and_persisted(store):
    sender = Mock(return_value='gmail-id')
    for index in range(2):
        job_id = str(index)
        digest = prepared(store, job_id=job_id)
        store.send_approved('alice', job_id, digest, {}, sender=sender)
    digest = prepared(store, job_id='third')
    with pytest.raises(ValueError, match='limit'):
        store.send_approved('alice', 'third', digest, {}, sender=sender)
    digest = prepared(store, tenant='bob')
    store.send_approved('bob', 'job1', digest, {}, sender=sender)
    assert sender.call_count == 3


def test_resume_change_invalidates_approval(store):
    digest = prepared(store)
    store.save_profile('alice', resume_text='New evidence')
    sender = Mock()
    with pytest.raises(ValueError, match='changed'):
        store.send_approved('alice', 'job1', digest, {}, sender=sender)
    sender.assert_not_called()


def test_remember_preserves_application_status(store):
    digest = prepared(store)
    store.send_approved('alice', 'job1', digest, {}, sender=Mock(return_value='sent'))
    store.remember('alice', [dict(job_id='job1', title='Updated title')])
    assert store.list_jobs('alice')[0]['status'] == 'awaiting_response'
    with pytest.raises(ValueError):
        store.track('alice', 'job1', 'needs_review')


def test_evidence_is_literal_and_does_not_invent_experience():
    text = 'MSc Finance at University\nIntern, Example 2024\nPython and SQL'
    profile = evidence_profile(text)
    assert profile['education_evidence'] == ['MSc Finance at University']
    assert profile['experience_evidence'] == ['Intern, Example 2024']
    assert profile['skills'] == ['python', 'sql']
    assert 'experience_years' not in profile


def test_provider_send_uses_https_without_retry():
    session = Mock()
    token, sent = Mock(), Mock()
    token.json.return_value = {'access_token': 'fake-test-token'}
    sent.json.return_value = {'id': 'test-message'}
    session.post.side_effect = [token, sent]
    assert send_gmail(dict(client_id='fake', client_secret='fake', refresh_token='fake', sender_email='alice@example.org'),
        'jobs@example.org', 'Subject', 'Body', b'%PDF-test', session) == 'test-message'
    assert session.post.call_count == 2
    assert session.post.call_args_list[1].args[0] == 'https://gmail.googleapis.com/gmail/v1/users/me/messages/send'
    with pytest.raises(ValueError):
        send_gmail({}, 'a@example.org,b@example.org', 'Subject', 'Body', b'', session)
    assert session.post.call_count == 2


def test_local_keyword_filter_and_feed_failure_isolation():
    session = Mock()
    himalayas, remotive = Mock(), Mock()
    himalayas.json.return_value = {'jobs': [dict(title='Unrelated writer', location='India', applicationLink='https://example.org/1'),
        dict(title='Financial analyst', location='India', applicationLink='https://example.org/2', description='Financial analyst') ]}
    remotive.raise_for_status.side_effect = ValueError('invalid provider response')
    session.get.side_effect = [himalayas, remotive]
    jobs, errors = fetch_public('financial analyst', {}, session=session)
    assert len(jobs) == 1 and jobs[0]['title'] == 'Financial analyst'
    assert len(errors) == 1 and errors[0].startswith('Remotive:')
    assert location_assessment('Remote') == 'Eligibility unverified'
    assert location_assessment('Indiana, USA') == 'Other region listed'


def test_hosted_app_does_not_expose_data_without_configuration(monkeypatch):
    monkeypatch.delenv('AUTOJOB_LOCAL_DEV', raising=False)
    from streamlit.testing.v1 import AppTest
    app = AppTest.from_file(Path(__file__).resolve().parents[1] / 'streamlit_app.py').run(timeout=20)
    assert not app.exception
    assert 'setup' in app.title[0].value.lower()
    assert not app.get('file_uploader')


def test_all_pages_render_in_explicit_local_mode(monkeypatch, tmp_path):
    monkeypatch.setenv('AUTOJOB_LOCAL_DEV', '1')
    monkeypatch.setenv('AUTOJOB_DATABASE_URL', 'sqlite:///' + str(tmp_path / 'ui.db'))
    from streamlit.testing.v1 import AppTest
    app = AppTest.from_file(Path(__file__).resolve().parents[1] / 'streamlit_app.py').run(timeout=20)
    for page in ['Overview', 'My loops', 'Board', 'All Matches', 'Search Jobs', 'My Applications', 'Questions', 'CV Checker', 'CV Builder', 'Inbox', 'Personal Info']:
        app.sidebar.radio[0].set_value(page).run(timeout=20)
        assert not app.exception, (page, app.exception)


@pytest.mark.skipif(not os.getenv('AUTOJOB_TEST_POSTGRES_URL'), reason='Dedicated PostgreSQL test database not supplied')
def test_postgres_concurrent_approval_is_single_send():
    import uuid
    value = Store(os.environ['AUTOJOB_TEST_POSTGRES_URL'])
    tenant = 'test-' + uuid.uuid4().hex
    value.ensure_user(tenant)
    value.save_profile(tenant, resume_text='Verified resume evidence', profile='{}')
    digest = prepared(value, tenant=tenant)
    sender = Mock(return_value='test-message')
    def attempt():
        try:
            return value.send_approved(tenant, 'job1', digest, {}, sender=sender)
        except ValueError:
            return 'blocked'
    try:
        with ThreadPoolExecutor(max_workers=2) as executor:
            results = list(executor.map(lambda _: attempt(), range(2)))
        assert sorted(results) == ['blocked', 'test-message']
        sender.assert_called_once()
    finally:
        with value.engine.begin() as conn:
            conn.execute(value.jobs.delete().where(value.jobs.c.tenant == tenant))
            conn.execute(value.users.delete().where(value.users.c.tenant == tenant))


def test_scheduler_only_discovers_opted_in_users(store):
    from scheduled_discovery import run
    store.save_profile('alice', profile=json.dumps({'agent_enabled': True, 'query': 'analyst', 'location': 'India only'}))
    fetcher = Mock(return_value=([dict(job_id='scheduled',title='Analyst',company='Example',location='India',description='Verified resume evidence')], []))
    assert run(store, fetcher) == 0
    fetcher.assert_called_once()
    assert store.list_jobs('alice')[0]['status'] == 'discovered'
    assert store.list_jobs('bob') == []


@pytest.mark.skipif(not os.getenv('AUTOJOB_TEST_POSTGRES_URL'), reason='Dedicated PostgreSQL test database not supplied')
def test_gmail_oauth_is_user_bound_encrypted_and_one_use():
    import uuid
    from urllib.parse import urlparse, parse_qs
    from cryptography.fernet import Fernet
    from gmail_connection import GmailConnections
    value = Store(os.environ['AUTOJOB_TEST_POSTGRES_URL'])
    tenant = 'oauth-test-' + uuid.uuid4().hex
    other = 'other-' + uuid.uuid4().hex
    connections = GmailConnections(value.engine, dict(client_id='fake-client', client_secret='fake-secret', redirect_uri='https://example.org/'), Fernet.generate_key().decode())
    try:
        url = connections.authorize_url(tenant, 'alice@example.org')
        state = parse_qs(urlparse(url).query)['state'][0]
        session = Mock()
        response = Mock()
        response.json.return_value = dict(id_token='fake-id', refresh_token='fake-refresh', scope='openid email https://www.googleapis.com/auth/gmail.send')
        session.post.return_value = response
        verifier = Mock(return_value=dict(email='alice@example.org', email_verified=True))
        with pytest.raises(ValueError, match='state'):
            connections.finish(other, 'alice@example.org', state, 'fake-code', session, verifier)
        session.post.assert_not_called()
        connections.finish(tenant, 'alice@example.org', state, 'fake-code', session, verifier)
        assert connections.credentials(tenant)['refresh_token'] == 'fake-refresh'
        assert connections.credentials(other) is None
        with value.engine.connect() as conn:
            stored = conn.execute(select(connections.table.c.token).where(connections.table.c.tenant == tenant)).scalar_one()
        assert stored != 'fake-refresh'
        with pytest.raises(ValueError, match='state'):
            connections.finish(tenant, 'alice@example.org', state, 'fake-code', session, verifier)
        session.post.assert_called_once()
        connections.disconnect(tenant)
        assert connections.credentials(tenant) is None
        state = parse_qs(urlparse(connections.authorize_url(tenant, 'alice@example.org')).query)['state'][0]
        verifier.return_value = dict(email='bob@example.org', email_verified=True)
        with pytest.raises(ValueError, match='same verified email'):
            connections.finish(tenant, 'alice@example.org', state, 'fake-code', session, verifier)
        assert connections.credentials(tenant) is None
    finally:
        with value.engine.begin() as conn:
            conn.execute(connections.table.delete().where(connections.table.c.tenant.in_([tenant, other])))


def test_prepared_application_ui_requires_final_approval(monkeypatch, tmp_path):
    url = 'sqlite:///' + str(tmp_path / 'review-ui.db')
    value = Store(url)
    value.ensure_user('local-development')
    value.save_profile('local-development', resume_text='Verified resume evidence', filename='resume.pdf', profile=json.dumps({'name':'Test Candidate','threshold':80}))
    prepared(value, tenant='local-development')
    monkeypatch.setenv('AUTOJOB_LOCAL_DEV', '1')
    monkeypatch.setenv('AUTOJOB_DATABASE_URL', url)
    from streamlit.testing.v1 import AppTest
    app = AppTest.from_file(Path(__file__).resolve().parents[1] / 'streamlit_app.py').run(timeout=20)
    app.sidebar.radio[0].set_value('My Applications').run(timeout=20)
    assert not app.exception
    send = next(b for b in app.button if b.label == 'Approve and send this email')
    assert send.disabled
    assert any('exact email' in c.label for c in app.checkbox)
    assert any('Email sending is disabled in local preview' in item.value for item in app.info)


def test_loop_wizard_saves_real_private_configuration(monkeypatch, tmp_path):
    url = 'sqlite:///' + str(tmp_path / 'wizard.db')
    monkeypatch.setenv('AUTOJOB_LOCAL_DEV', '1')
    monkeypatch.setenv('AUTOJOB_DATABASE_URL', url)
    from streamlit.testing.v1 import AppTest
    app = AppTest.from_file(Path(__file__).resolve().parents[1] / 'streamlit_app.py').run(timeout=20)
    def field(kind, label):
        return next(w for w in getattr(app, kind) if w.label == label)
    field('button', 'Continue').click().run()
    assert app.error and app.session_state.wizard_step == 0
    field('text_input', 'Job Titles').set_value('Financial analyst')
    field('multiselect', 'Search in Specific Job Boards').set_value(['Remotive'])
    field('button', 'Continue').click().run()
    assert app.session_state.wizard_step == 1
    field('text_input', 'Email Subject').set_value('Application — {job_title}')
    field('button', 'Continue').click().run()
    assert app.session_state.wizard_step == 2
    field('text_input', 'Loop Name').set_value('My finance search')
    field('number_input', 'Minimum application match (%)').set_value(90)
    field('button', 'Continue').click().run()
    assert app.session_state.wizard_step == 3
    field('button', 'Create Loop').click().run()
    assert not app.exception
    assert app.session_state.active_page == 'My loops'
    value = Store(url)
    rows = value.list_loops('local-development')
    assert len(rows) == 1 and rows[0]['enabled'] is False
    config = json.loads(rows[0]['config'])
    assert config['name'] == 'My finance search' and config['threshold'] == 90
    assert config['sources'] == ['Remotive']
    assert value.list_loops('another-user') == []
    value.set_loop_enabled('another-user', rows[0]['loop_id'], True)
    assert value.list_loops('local-development')[0]['enabled'] is False


def test_selected_sources_do_not_call_other_providers():
    session = Mock()
    response = Mock()
    response.json.return_value = {'jobs': []}
    session.get.return_value = response
    jobs, errors = fetch_public('analyst', session=session, enabled_sources=['Remotive'])
    assert jobs == [] and errors == []
    session.get.assert_called_once()
    assert session.get.call_args.args[0] == 'https://remotive.com/api/remote-jobs'


def test_email_templates_cannot_invent_nested_fields():
    from loops import validate_template, email_preview
    for template in ['{name.__class__}', '{name[0]}', '{unsupported}', '{name!r}', '{name:>10000000}', '{']:
        with pytest.raises(ValueError):
            validate_template(template)
    subject, body = email_preview({}, dict(title='Analyst',company='Example',url='https://example.org'), 'Candidate')
    assert 'Candidate' in subject and 'Analyst' in body
    assert 'years' not in body


def test_scheduler_runs_enabled_saved_loops_without_sending(store):
    from scheduled_discovery import run
    from loops import LOCATIONS
    config = dict(query='analyst', location=LOCATIONS[0], sources=['Remotive'])
    loop_id = store.save_loop('alice', config, enabled=True)
    store.save_loop('alice', dict(config, query='engineer'), enabled=False)
    fetcher = Mock(return_value=([dict(job_id='loop-job', title='Analyst', company='Example', location='India', description='Analyst')], []))
    assert run(store, fetcher) == 0
    fetcher.assert_called_once()
    assert fetcher.call_args.kwargs['enabled_sources'] == ['Remotive']
    row = store.list_jobs('alice')[0]
    assert row['status'] == 'discovered'
    assert json.loads(row['payload'])['loop_id'] == loop_id
    assert store.list_jobs('bob') == []


def test_loop_threshold_is_checked_again_at_dispatch(store):
    digest = prepared(store)
    row = store.job('alice', 'job1')
    app = json.loads(row['application'])
    app['loop_threshold'] = 95
    store.prepare('alice', 'job1', app, bytes(row['attachment']))
    sender = Mock()
    with pytest.raises(ValueError, match='threshold'):
        store.send_approved('alice', 'job1', digest, {}, sender=sender)
    sender.assert_not_called()


def test_docx_resume_includes_tables_in_document_order():
    import io
    from docx import Document
    from resume_match import read_resume
    document = Document()
    document.add_paragraph('Candidate résumé')
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = 'Python and SQL'
    table.cell(0, 1).text = 'MSc at Example University'
    table.cell(1, 0).merge(table.cell(1, 1)).text = 'Intern at Example, 2024'
    document.add_paragraph('End of résumé')
    raw = io.BytesIO()
    document.save(raw)
    extracted = read_resume('resume.docx', raw.getvalue())
    assert extracted.splitlines() == ['Candidate résumé', 'Python and SQL', 'MSc at Example University', 'Intern at Example, 2024', 'End of résumé']
    profile = evidence_profile(extracted)
    assert profile['skills'] == ['python', 'sql']
    assert profile['education_evidence'] == ['MSc at Example University']


@pytest.mark.parametrize('location,worldwide,expected', [
    ('Worldwide remote, US only', False, 'Eligibility unverified'),
    ('Canada', True, 'Eligibility unverified'),
    ('Worldwide except India', False, 'Other region listed'),
    ('India not eligible', False, 'Other region listed'),
    ('Remote, India, US only', False, 'Eligibility unverified'),
    ('India or United States', False, 'India listed'),
    ('Worldwide remote', False, 'Worldwide listed'),
])
def test_location_metadata_does_not_override_restrictions(location, worldwide, expected):
    assert location_assessment(location, worldwide) == expected


def test_missing_feed_location_is_not_invented_as_worldwide():
    session, response = Mock(), Mock()
    response.json.return_value = {'jobs': [dict(id='missing-location', title='Python developer', company_name='Example', url='https://example.org/job', description='Python') ]}
    session.get.return_value = response
    jobs, errors = fetch_public('python', session=session, enabled_sources=['Remotive'])
    assert jobs == [] and errors == []
    jobs, errors = fetch_public('python', session=session, enabled_sources=['Remotive'], location_policy='all')
    assert len(jobs) == 1 and errors == []
    assert jobs[0]['location'] == 'Location not provided'
    assert jobs[0]['location_status'] == 'Eligibility unverified'


def test_matching_tolerates_null_optional_provider_fields():
    from resume_match import score
    match = score('Python and SQL', {'title': 'Developer', 'description': None})
    assert 0 <= match['match_score'] <= 100
    assert match['eligibility'] == 'Review required'


def test_indian_degree_evidence_does_not_become_work_experience():
    profile = evidence_profile('B.Tech Computer Science, Example Institute, 2024\nMCA, Example College\nDeveloper internship, Example Company, 2025')
    assert profile['education_evidence'] == ['B.Tech Computer Science, Example Institute, 2024', 'MCA, Example College']
    assert profile['experience_evidence'] == ['Developer internship, Example Company, 2025']
    assert 'experience_years' not in profile


def test_accented_resume_text_does_not_invent_r_programming():
    from resume_match import skills_in
    assert skills_in('My résumé: Python and SQL') == {'python', 'sql'}
    assert skills_in('R programming; Python; SQL') == {'r', 'python', 'sql'}
