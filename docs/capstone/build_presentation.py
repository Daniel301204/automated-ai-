"""Create the editable capstone deck and a matching PDF. No private data used.

Run from the repository:
  .venv/bin/python -m pip install -r docs/capstone/requirements.txt
  .venv/bin/python docs/capstone/build_presentation.py
"""
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt
from reportlab.lib.colors import HexColor
from reportlab.pdfbase.pdfmetrics import registerFont, stringWidth
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parent
WIDTH, HEIGHT = 13.333, 7.5
NAVY, BLUE, INK, MUTED = "13253F", "1967C3", "172B46", "5C6D82"
PALE, WHITE, GREEN, AMBER = "F1F5FA", "FFFFFF", "11755F", "936510"
PDF_REGULAR, PDF_BOLD = 'Helvetica', 'Helvetica-Bold'
for regular, bold in [
    ('/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf', '/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf'),
    ('C:/Windows/Fonts/arial.ttf', 'C:/Windows/Fonts/arialbd.ttf'),
]:
    if Path(regular).is_file() and Path(bold).is_file():
        registerFont(TTFont('LC-Regular', regular))
        registerFont(TTFont('LC-Bold', bold))
        PDF_REGULAR, PDF_BOLD = 'LC-Regular', 'LC-Bold'
        break


class Deck:
    def __init__(self):
        self.ppt = Presentation()
        self.ppt.slide_width, self.ppt.slide_height = Inches(WIDTH), Inches(HEIGHT)
        self.ppt.core_properties.title = "LoopingCV — Capstone Project"
        self.ppt.core_properties.subject = "Explainable job discovery and approved applications"
        self.ppt.core_properties.author = "LoopingCV project"
        self.pdf = canvas.Canvas(str(ROOT / "LoopingCV_Capstone.pdf"),
                                 pagesize=(WIDTH * 72, HEIGHT * 72))
        self.pdf.setTitle("LoopingCV — Capstone Project")
        self.notes = []

    def slide(self, section, title, subtitle="", *, dark=False, notes=""):
        if self.ppt.slides:
            self.pdf.showPage()
        self.current = self.ppt.slides.add_slide(self.ppt.slide_layouts[6])
        self.current.background.fill.solid()
        self.current.background.fill.fore_color.rgb = RGBColor.from_string(NAVY if dark else WHITE)
        self.pdf.setFillColor(HexColor("#" + (NAVY if dark else WHITE)))
        self.pdf.rect(0, 0, WIDTH * 72, HEIGHT * 72, fill=1, stroke=0)
        self.box(0, 0, .14, HEIGHT, BLUE, radius=False)
        self.text(section.upper(), .62, .36, 11.7, .3, 11, "94BEFF" if dark else BLUE, bold=True)
        if title:
            self.text(title, .62, .88, 12.0, .61, 30, WHITE if dark else INK, bold=True)
        if subtitle:
            self.text(subtitle, .65, 1.52, 11.9, .43, 15, "CEDCEF" if dark else MUTED)
        number = len(self.ppt.slides)
        self.text("LoopingCV  |  Capstone prototype", .65, 7.05, 10, .23, 9,
                  "B2C4DE" if dark else MUTED)
        self.text(f"{number:02d} / 12", 12, 7.05, .7, .23, 9,
                  "B2C4DE" if dark else MUTED)
        self.current.notes_slide.notes_text_frame.text = notes
        self.notes.append((title or "LoopingCV", notes))

    def text(self, content, x, y, width, height, size=16, color=INK, bold=False):
        font = PDF_BOLD if bold else PDF_REGULAR
        lines = []
        for raw in str(content).split("\n"):
            line = ""
            for word in raw.split():
                candidate = (line + " " + word).strip()
                if line and stringWidth(candidate, font, size) > width * 72 - 3:
                    lines.append(line)
                    line = word
                else:
                    line = candidate
            lines.append(line)
        leading = size * 1.22
        if len(lines) * leading > height * 72 + 1:
            raise ValueError(f"Text exceeds slide box: {content!r}")
        shape = self.current.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
        frame = shape.text_frame
        frame.word_wrap = False
        frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
        for index, line in enumerate(lines):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.text = line
            paragraph.font.name = "Arial"
            paragraph.font.size = Pt(size)
            paragraph.font.bold = bold
            paragraph.font.color.rgb = RGBColor.from_string(color)
            paragraph.space_before = paragraph.space_after = Pt(0)
            paragraph.line_spacing = Pt(leading)
        self.pdf.setFillColor(HexColor("#" + color))
        self.pdf.setFont(font, size)
        for index, line in enumerate(lines):
            self.pdf.drawString(x * 72, (HEIGHT - y) * 72 - size - index * leading, line)

    def box(self, x, y, width, height, color=PALE, *, radius=True):
        shape = self.current.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
            Inches(x), Inches(y), Inches(width), Inches(height))
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(color)
        shape.line.fill.background()
        if radius:
            shape.adjustments[0] = .08
        self.pdf.setFillColor(HexColor("#" + color))
        if radius:
            self.pdf.roundRect(x * 72, (HEIGHT - y - height) * 72,
                               width * 72, height * 72, 7, fill=1, stroke=0)
        else:
            self.pdf.rect(x * 72, (HEIGHT - y - height) * 72,
                          width * 72, height * 72, fill=1, stroke=0)

    def arrow(self, x1, y1, x2, y2):
        line = self.current.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        line.line.color.rgb = RGBColor.from_string(BLUE)
        line.line.width = Pt(1.8)
        end = OxmlElement("a:tailEnd")
        end.set("type", "triangle")
        line._element.spPr.get_or_add_ln().append(end)
        self.pdf.setStrokeColor(HexColor("#" + BLUE))
        self.pdf.setLineWidth(1.8)
        self.pdf.line(x1 * 72, (HEIGHT - y1) * 72, x2 * 72, (HEIGHT - y2) * 72)

    def card(self, x, y, width, height, heading, body, *, tag="", color=PALE):
        self.box(x, y, width, height, color)
        offset = .18
        if tag:
            self.text(tag.upper(), x + .2, y + offset, width - .4, .3, 10, BLUE, bold=True)
            offset += .38
        self.text(heading, x + .2, y + offset, width - .4, .72, 20, bold=True)
        self.text(body, x + .2, y + offset + .83, width - .4,
                  height - offset - 1.02, 15, MUTED)

    def image(self, path, x, y, width, height):
        with Image.open(path) as image:
            iw, ih = image.size
        scale = min(width / iw, height / ih)
        w, h = iw * scale, ih * scale
        x += (width - w) / 2
        y += (height - h) / 2
        self.current.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
        self.pdf.drawImage(str(path), x * 72, (HEIGHT - y - h) * 72, w * 72, h * 72)

    def save(self):
        assert len(self.ppt.slides) == 12
        self.ppt.save(ROOT / "LoopingCV_Capstone.pptx")
        self.pdf.save()
        notes = ["# LoopingCV: presenter notes", "", "Replace the student and department placeholders on the cover.", ""]
        for number, (title, note) in enumerate(self.notes, 1):
            notes += [f"## Slide {number}: {title}", "", note, ""]
        (ROOT / "PRESENTER_NOTES.md").write_text("\n".join(notes), encoding="utf-8")


def main():
    deck = Deck()
    deck.slide("Capstone project", "", dark=True,
               notes="Introduce LoopingCV as an explainable job-search and application-review prototype. Replace the student and department placeholders. This is not a claim that the public service or every requested integration is already complete.")
    deck.text("LoopingCV", .75, 1.35, 11.5, 1.15, 58, WHITE, bold=True)
    deck.text("Evidence-based job search.\nApplications you review and approve.", .78, 2.72, 11.5, 1.2, 28, "D6E5FA")
    deck.box(.8, 4.4, 5.6, .65, BLUE)
    deck.text("Public URL created  |  Account setup pending", 1.0, 4.58, 5.25, .32, 13, WHITE)
    deck.text("Student name: ____________________\nRegister number / Department: ____________________", .8, 5.45, 11.5, .85, 17, "D6E5FA")

    deck.slide("01 / Problem and objective", "Reduce repetitive job-search work",
               "Keep the user in control of facts, eligibility and every outgoing application.",
               notes="The problem is repetitive discovery, repeated résumé uploads and scattered tracking. The objective is a private workspace that centralizes the workflow while requiring explicit human verification. It does not guarantee that a candidate qualifies or receives interviews.")
    deck.card(.65, 2.13, 5.8, 3.65, "Current difficulties",
              "Searching many job sites manually\nRe-entering résumé information\nWriting generic application messages\nTracking responses across different tools", tag="Problem")
    deck.card(6.85, 2.13, 5.8, 3.65, "LoopingCV approach",
              "Upload once and review evidence\nSave India / remote search loops\nRank actual public listings\nApprove exact messages and track status", tag="Objective")
    deck.text("Design principle: uncertainty becomes a review task.", .75, 6.24, 11.8, .4, 19, BLUE, bold=True)

    deck.slide("02 / Workflow", "From résumé to reviewed application",
               "Discovery can run on a schedule; sending requires a user's explicit approval.",
               notes="Explain the six stages. Extraction reads literal evidence and preferences are entered by the user. Listings come from actual public feeds. A match score is only a similarity signal. Gmail transport is implemented but needs configuration; automated ATS form submission is not implemented.")
    steps = [
        ("01", "Upload résumé", "Read PDF, DOCX or TXT"),
        ("02", "Review profile", "Confirm facts and preferences"),
        ("03", "Create a loop", "Choose sources, region and limits"),
        ("04", "Discover and rank", "Actual listings; explainable scores"),
        ("05", "Verify and approve", "Employer contact, eligibility and PDF"),
        ("06", "Track applications", "Dispatch state and confirmed updates"),
    ]
    for index, (number, title, body) in enumerate(steps):
        x, y = .65 + (index % 3) * 4.1, 2.15 + (index // 3) * 1.83
        deck.box(x, y, 3.85, 1.5)
        deck.text(number, x + .18, y + .16, .6, .34, 17, BLUE, bold=True)
        deck.text(title, x + .18, y + .59, 3.45, .4, 18, bold=True)
        deck.text(body, x + .18, y + 1.03, 3.45, .43, 13, MUTED)
    deck.text("ATS applications currently open the official employer page for manual completion.", .75, 6.28, 11.8, .4, 16, MUTED)

    deck.slide("03 / Architecture", "A private workspace with external integrations",
               "Python + Streamlit + SQLAlchemy + PostgreSQL + per-user Google OAuth.",
               notes="The UI uses Google OIDC for private sign-in. The workflow layer runs extraction, scoring, review and approval checks. PostgreSQL isolates records by a stable identity. Public feeds are read-only. Gmail uses a user's own encrypted refresh token and only sends after a database reservation. The scheduler discovers jobs and never sends applications.")
    deck.card(.65, 2.14, 3.75, 2.55, "Streamlit interface",
              "Google sign-in\nRésumé and loop screens\nReview and tracking", tag="User")
    deck.card(4.79, 2.14, 3.75, 2.55, "Workflow engine",
              "Literal extraction and scoring\nVerification and exact approval\nQuota / duplicate checks", tag="Python")
    deck.card(8.94, 2.14, 3.75, 2.55, "Private PostgreSQL",
              "Tenant-separated records\nOriginal résumé and drafts\nEncrypted Gmail connection", tag="Storage")
    deck.arrow(4.4, 3.2, 4.76, 3.2)
    deck.arrow(8.54, 3.2, 8.91, 3.2)
    deck.box(.65, 4.85, 12.04, 1.22, NAVY)
    deck.text("Public feeds: Himalayas / Remotive / Greenhouse / Lever", .92, 5.07, 11.4, .33, 17, WHITE)
    deck.text("External scheduler: discovery only   |   Gmail API: individually approved emails", .92, 5.53, 11.4, .32, 15, "D6E5FA")

    deck.slide("04 / Matching method", "Explainable keyword matching",
               "A transparent scoring baseline. No trained language model or LLM API is required.",
               notes="Explain the actual implementation: recognized skill overlap contributes 50%, scaled keyword overlap 32% and title overlap 18%. If a job has no recognized skills, keyword overlap substitutes for that component. Scores are clamped to 0–100. These are heuristic similarity scores, not calibrated probabilities. Accuracy has not been measured on a labeled dataset.")
    for x, number, heading, description in [
        (.65, "50%", "Skill overlap", "Compare recognized skills in the résumé and job text."),
        (4.79, "32%", "Keyword overlap", "Compare normalized job and résumé words."),
        (8.94, "18%", "Title overlap", "Compare the role title against résumé terms."),
    ]:
        deck.box(x, 2.17, 3.75, 2.72)
        deck.text(number, x + .23, 2.4, 3.2, .84, 40, BLUE, bold=True)
        deck.text(heading, x + .23, 3.37, 3.2, .43, 20, bold=True)
        deck.text(description, x + .23, 4.0, 3.2, .75, 15, MUTED)
    deck.text("Output: matched skills, missing skills, shared keywords and a similarity score.", .72, 5.35, 11.9, .65, 20, bold=True)
    deck.text("Eligibility, work authorization and mandatory requirements still need human verification.", .72, 6.24, 11.9, .42, 15, MUTED)

    deck.slide("05 / Application interface", "LoopingCV dashboard",
               "The actual local interface: a four-step search-loop wizard and private navigation.",
               notes="This screenshot is an actual browser render of the current local interface. It is not a deployed public website or evidence of a completed job application. Show the New Loop wizard, CV Checker, All Matches, My Applications and Board. Share app becomes useful after a valid public deployment is configured.")
    deck.box(.65, 2.07, 8.25, 4.61)
    deck.image(ROOT / "dashboard.png", .72, 2.12, 8.1, 4.5)
    deck.text("One place to work", 9.18, 2.35, 3.45, .48, 22, bold=True)
    deck.text("CV Checker\nNew Loop\nSearch Jobs / All Matches\nMy Applications\nBoard and CSV export\nShare app", 9.18, 3.13, 3.45, 2.83, 19, MUTED)
    deck.text("Screenshot: local preview", 9.18, 6.27, 3.45, .34, 12, BLUE)

    deck.slide("06 / Reliability", "Safeguards around each outgoing email",
               "Approval is checked against the recipient, subject, message and original PDF.",
               notes="Discuss why approved sending needs PostgreSQL concurrency protection. Approval fingerprints prevent changing an already-approved message. Per-user locking enforces quotas. Sent or uncertain applications cannot be automatically retried. Remote database TLS is verified and Gmail tokens are encrypted. These safeguards are covered by tests but real delivery still needs a deployed acceptance test.")
    safeguards = [
        ("Exact approval", "A message/PDF fingerprint must match the saved draft."),
        ("Private accounts", "Records and Gmail credentials are scoped to each user."),
        ("Daily attempt limits", "Database locks protect per-user quota reservations."),
        ("Duplicate protection", "An already-attempted application cannot be resent."),
        ("Uncertain delivery", "Timeouts require reconciliation in Gmail Sent."),
        ("Credential protection", "Verified database TLS and encrypted OAuth tokens."),
    ]
    for index, (title, body) in enumerate(safeguards):
        x, y = .65 + (index % 3) * 4.1, 2.14 + (index // 3) * 2.06
        deck.box(x, y, 3.85, 1.84)
        deck.text(title, x + .2, y + .2, 3.45, .58, 19, bold=True)
        deck.text(body, x + .2, y + .97, 3.45, .75, 15, MUTED)

    deck.slide("07 / Validation", "32 automated checks passed",
               "Validated on 11 October 2026, including a real temporary PostgreSQL database.",
               notes="The full suite passed with 32 checks and no skips, including concurrent approvals and encrypted, user-bound, one-use OAuth persistence. Provider sends in the suite are mocked: these tests do not establish real delivery. Browser checks exercised desktop/mobile navigation. Live read-only Greenhouse discovery also returned actual India listings. Model accuracy was not measured.")
    deck.box(.65, 2.18, 3.05, 3.83, BLUE)
    deck.text("32", .93, 2.5, 2.5, 1.22, 65, WHITE, bold=True)
    deck.text("passed\n0 failed\n0 skipped", .95, 4.0, 2.45, 1.18, 23, WHITE)
    deck.text("Confirmed fixes", 4.12, 2.3, 8.25, .49, 23, bold=True)
    deck.text("DOCX table contents are included.\nAccented “résumé” no longer invents an R skill.\nMissing or conflicting locations remain uncertain.\nDegree dates are not treated as work experience.\nNull optional job text does not crash matching.", 4.12, 3.12, 8.25, 2.23, 18, MUTED)
    deck.text("Tests verify software behavior; they are not a model-accuracy score.", 4.12, 5.76, 8.25, .66, 16, BLUE, bold=True)

    deck.slide("08 / Honest feature status", "Implemented behavior and remaining work",
               "The capstone scope separates working code from configuration and future integrations.",
               notes="Be candid about status. Résumé extraction, loops, discovery, matching, review and tracking are implemented. Gmail transport and scheduling need external configuration. ATS submission and inbox synchronization are not implemented. Interviews and rejections are manually confirmed. The public deployment needs the user's Streamlit, database and Google accounts.")
    rows = [
        ("Résumé / loops / discovery / matching", "Implemented"),
        ("Exact draft approval / tracking / CSV", "Implemented"),
        ("Per-user Gmail sending", "Built; OAuth setup + live delivery test pending"),
        ("Scheduled discovery", "Built; Actions secret + default-branch setup"),
        ("Public URL / WhatsApp app sharing", "URL created; private configuration pending"),
        ("Automatic ATS forms / inbox replies", "Not implemented; manual confirmation today"),
        ("OCR / trained semantic model", "Not implemented"),
    ]
    for index, (feature, state) in enumerate(rows):
        y = 2.17 + index * .58
        deck.box(.65, y, 12.02, .52, PALE if index % 2 == 0 else WHITE, radius=False)
        deck.text(feature, .84, y + .14, 5.3, .3, 14, bold=True)
        deck.text(state, 6.26, y + .14, 6.15, .3, 14, MUTED)
    deck.text("No submission or employer response is claimed without a confirmed action.", .78, 6.58, 11.8, .28, 13, BLUE)

    deck.slide("09 / Public deployment", "Finish configuring the public app",
               "Streamlit Community Cloud hosts the app; Google and PostgreSQL enable private use.",
               notes="The setup screen means the app code is running but auth/database secrets are missing. Use the actual Streamlit app URL for both Google callbacks. Enter private credentials in Streamlit Secrets. Add demo friends to Google test users while OAuth is in testing. Validate separate accounts and approved Gmail delivery before presenting the integration as live.")
    deck.card(.65, 2.14, 3.75, 3.68, "App URL created",
              "Repository: automated-ai-\nBranch: codex/approval-gated-streamlit\nFile: streamlit_app.py\nPython: 3.12", tag="Step 1")
    deck.card(4.79, 2.14, 3.75, 3.68, "Configure secrets",
              "Managed PostgreSQL\nGoogle Web OAuth + Gmail API\nCookie and encryption keys\nActual app callback URLs", tag="Step 2")
    deck.card(8.94, 2.14, 3.75, 3.68, "Validate and share",
              "Two private test accounts\nReviewed résumé and draft\nUser-approved test email\nShare app → WhatsApp", tag="Step 3")
    deck.text("Public app: https://5dasesqq5hdkocoz5dupwp.streamlit.app/", .75, 6.15, 11.8, .38, 15, BLUE)
    deck.text("The app currently shows setup. Add private secrets before inviting demo users.", .75, 6.6, 11.8, .32, 13, MUTED)

    deck.slide("10 / Capstone demonstration", "A short, reproducible demonstration",
               "Demonstrate the actual state; use clearly labeled demo data when presenting.",
               notes="Walk through these six actions. In local mode, sending is disabled and the presentation should show the saved draft and attachment review. Once hosted sign-in, PostgreSQL and per-user Gmail are configured, the owner can approve an intentional real test email and verify Gmail Sent. Do not send applications using fictional demo credentials or a fictional résumé.")
    demo_steps = [
        "Upload a readable résumé and review the extracted evidence.",
        "Create an India / remote loop with a reviewed template and limits.",
        "Discover real vacancies and inspect matched / missing skills.",
        "Verify mandatory criteria and the employer-published contact.",
        "Review the exact draft and original PDF before any approval.",
        "Inspect Board, export CSV and share the configured public link.",
    ]
    for index, content in enumerate(demo_steps):
        y = 2.14 + index * .62
        deck.box(.72, y, .42, .4, BLUE)
        deck.text(str(index + 1), .84, y + .07, .25, .25, 12, WHITE, bold=True)
        deck.text(content, 1.42, y + .03, 11.1, .44, 18, INK)
    deck.text("Local demo: no real email sends. Hosted integration demo: explicit approval + Gmail Sent check.", .77, 6.37, 11.8, .43, 14, BLUE)

    deck.slide("11 / Next steps", "A reliable baseline for further development",
               "Complete external configuration first, then evaluate each new integration.",
               notes="Conclude with measurable next steps rather than promising perfection. Publish and validate the authenticated app; evaluate ranking on a labeled dataset; add only authorized ATS provider adapters; implement inbox tracking with the required additional OAuth scopes and consent. The current version is a tested prototype with documented limits.")
    deck.card(.65, 2.18, 3.75, 3.25, "Evaluate matching",
              "Build a labeled résumé/job dataset.\nMeasure ranking quality and errors.\nCompare a semantic model with the keyword baseline.", tag="Future work")
    deck.card(4.79, 2.18, 3.75, 3.25, "Add ATS adapters",
              "Choose a supported provider API.\nObtain authorized integration access.\nKeep review, consent and truthful status.", tag="Future work")
    deck.card(8.94, 2.18, 3.75, 3.25, "Extend tracking",
              "Add consented inbox read access.\nClassify real replies conservatively.\nImprove monitoring and worker scale.", tag="Future work")
    deck.text("Today: explainable discovery, private review and approval safeguards.", .75, 5.83, 11.9, .52, 23, BLUE, bold=True)
    deck.text("Questions", .75, 6.48, 11.9, .35, 16, MUTED)
    deck.save()
    print("Created 12-slide PPTX, matching PDF and presenter notes.")


if __name__ == "__main__":
    main()
